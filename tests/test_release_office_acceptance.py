"""Stage 7 independent Office release acceptance on disposable synthetic data.

Exercise the *public engine* with real, richer documents, rather than only the
LibreOffice adapter. All fixtures and outputs are kept in pytest's tmp_path;
there are no UI, registry, package installation, or system-profile changes.
"""

from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path
import re
import zipfile

import pytest

from rightconvert.engine import convert_file, get_targets
from rightconvert.formats import ConversionError, inspect_source
from rightconvert.office_provider import find_libreoffice, inspect_office
from test_office_provider import _native_source, _create_odf


ROOT = Path(__file__).resolve().parents[1]
QA_SOFFICE = (
    ROOT / "build" / "stage6" / "libreoffice-isolated"
    / "portable-root" / "program" / "soffice.com"
)
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_ODF_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
_ODF_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_ODF_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
_ODF_DRAW = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
_ODF_SVG = "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _alter_archive(path: Path, changes: dict[str, bytes]) -> None:
    """Replace/add archive parts in memory without changing any installed file."""
    contents = io.BytesIO()
    with zipfile.ZipFile(path, "r") as original:
        with zipfile.ZipFile(contents, "w", compression=zipfile.ZIP_DEFLATED) as output:
            seen = set()
            for info in original.infolist():
                name = info.filename
                output.writestr(
                    name, changes.get(name, original.read(info)),
                    compress_type=zipfile.ZIP_STORED if name == "mimetype"
                    else zipfile.ZIP_DEFLATED,
                )
                seen.add(name)
            for name, data in changes.items():
                if name not in seen:
                    output.writestr(name, data)
    path.write_bytes(contents.getvalue())


def make_document(path: Path) -> tuple[str, tuple[str, ...], int]:
    """Make a richer passive document; return family, text markers, min pages."""
    kind = path.suffix.lstrip(".")
    markers = {
        "docx": ("Stage7 Word Statement", "Invoice Row Alpha", "Second Page Evidence"),
        "xlsx": ("Stage7 Spreadsheet", "Bucharest Revenue", "Operations 2026"),
        "pptx": ("Stage7 Presentation Intro", "Stage7 Presentation Results"),
        "odt": ("Stage7 ODT Section One", "Stage7 ODT Section Two"),
        "ods": ("Stage7 ODS Revenue", "Stage7 ODS Second Row"),
        "odp": ("Stage7 ODP Intro", "Stage7 ODP Summary"),
    }[kind]

    if kind == "docx":
        _native_source(path)
        w = _W
        data = (
            f'<w:document xmlns:w="{w}"><w:body>'
            f'<w:p><w:pPr><w:pStyle w:val="Title"/></w:pPr><w:r><w:t>{markers[0]}</w:t></w:r></w:p>'
            '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Bold receipt item</w:t></w:r></w:p>'
            '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="5000"/></w:tblGrid>'
            f'<w:tr><w:tc><w:p><w:r><w:t>{markers[1]}</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
            '<w:p><w:r><w:br w:type="page"/></w:r></w:p>'
            f'<w:p><w:r><w:t>{markers[2]}</w:t></w:r></w:p>'
            '<w:sectPr/></w:body></w:document>'
        ).encode()
        _alter_archive(path, {"word/document.xml": data})
        return kind, markers, 2

    if kind == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            pytest.skip("Realistic XLSX generation needs test-only openpyxl")
        wb = Workbook()
        first = wb.active
        first.title = "Revenue"
        first["A1"] = markers[0]
        first["A1"].font = Font(bold=True, size=18)
        first["A1"].fill = PatternFill(fill_type="solid", fgColor="E8EDF2")
        first["A2"] = markers[1]
        first["B2"] = 1865
        first["A3"] = "Total RON"
        first["B3"] = 7436
        first.column_dimensions["A"].width = 38
        first.column_dimensions["B"].width = 18
        second = wb.create_sheet("Operations")
        second["A1"] = markers[2]
        second["B1"] = "6 months"
        second.column_dimensions["A"].width = 36
        second.column_dimensions["B"].width = 18
        wb.save(path)
        return kind, markers, 2

    if kind == "pptx":
        _native_source(path)
        with zipfile.ZipFile(path, "r") as existing:
            slide1 = existing.read("ppt/slides/slide1.xml")
            presentation = existing.read("ppt/presentation.xml")
            rels = existing.read("ppt/_rels/presentation.xml.rels")
            types = existing.read("[Content_Types].xml")
        slide1 = slide1.replace(
            b"RightConvert native PPTX verification", markers[0].encode()
        )
        slide2 = slide1.replace(markers[0].encode(), markers[1].encode())
        presentation = presentation.replace(
            b"</p:sldIdLst>",
            b'<p:sldId id="257" r:id="rId2"/></p:sldIdLst>',
        )
        rels = rels.replace(
            b"</Relationships>",
            f'<Relationship Id="rId2" Type="{_R}/slide" '
            'Target="slides/slide2.xml"/></Relationships>'.encode(),
        )
        types = types.replace(
            b"</Types>",
            b'<Override PartName="/ppt/slides/slide2.xml" '
            b'ContentType="application/vnd.openxmlformats-officedocument.'
            b'presentationml.slide+xml"/></Types>',
        )
        _alter_archive(path, {
            "ppt/slides/slide1.xml": slide1,
            "ppt/slides/slide2.xml": slide2,
            "ppt/presentation.xml": presentation,
            "ppt/_rels/presentation.xml.rels": rels,
            "[Content_Types].xml": types,
        })
        return kind, markers, 2

    office = _ODF_OFFICE
    text = _ODF_TEXT
    table = _ODF_TABLE
    draw = _ODF_DRAW
    svg = _ODF_SVG
    ns = (
        f'xmlns:office="{office}" xmlns:text="{text}" '
        f'xmlns:table="{table}" xmlns:draw="{draw}" xmlns:svg="{svg}"'
    )
    mime = {
        "odt": "application/vnd.oasis.opendocument.text",
        "ods": "application/vnd.oasis.opendocument.spreadsheet",
        "odp": "application/vnd.oasis.opendocument.presentation",
    }[kind]
    if kind == "odt":
        body = (
            "<office:text>"
            f"<text:h text:outline-level=\"1\">{markers[0]}</text:h>"
            "<text:p>Stage7 standalone offline paragraphs and Unicode: București</text:p>"
            f"<text:p>{markers[1]}</text:p>"
            "</office:text>"
        )
        minimum = 1
    elif kind == "ods":
        body = (
            '<office:spreadsheet><table:table table:name="Revenue">'
            '<table:table-column/>'
            '<table:table-row><table:table-cell office:value-type="string">'
            f'<text:p>{markers[0]}</text:p></table:table-cell></table:table-row>'
            '<table:table-row><table:table-cell office:value-type="string">'
            f'<text:p>{markers[1]}</text:p></table:table-cell></table:table-row>'
            '</table:table></office:spreadsheet>'
        )
        minimum = 1
    else:
        def page(name: str, marker: str) -> str:
            return (
                f'<draw:page draw:name="{name}">'
                '<draw:frame svg:x="1cm" svg:y="1cm" svg:width="23cm" svg:height="4cm">'
                f'<draw:text-box><text:p>{marker}</text:p></draw:text-box>'
                '</draw:frame></draw:page>'
            )
        body = (
            '<office:presentation>' + page("Intro", markers[0])
            + page("Results", markers[1]) + '</office:presentation>'
        )
        minimum = 2
    content = (
        f'<office:document-content {ns} office:mimetype="{mime}">'
        f'<office:body>{body}</office:body></office:document-content>'
    ).encode()
    _create_odf(path, extra={"content.xml": content})
    return kind, markers, minimum


@pytest.fixture(scope="module")
def soffice() -> str:
    # Stage 7's default unit suite should stay independent from a QA-only
    # LibreOffice extraction and not launch six costly native subprocesses.
    # Run explicitly with RIGHTCONVERT_STAGE7_OFFICE_NATIVE=1.
    if os.environ.get("RIGHTCONVERT_STAGE7_OFFICE_NATIVE") != "1":
        pytest.skip("Native release Office QA opt-in: set RIGHTCONVERT_STAGE7_OFFICE_NATIVE=1")
    if not QA_SOFFICE.is_file():
        pytest.skip("Official LibreOffice QA extraction absent; native test not claimed")
    chosen = find_libreoffice(str(QA_SOFFICE))
    assert chosen == str(QA_SOFFICE.resolve(strict=True))
    return chosen


def _extract_pdf(pdf: Path) -> tuple[int, str]:
    renderer = pytest.importorskip("pypdfium2", reason="Independent PDFium parser unavailable")
    with renderer.PdfDocument(str(pdf)) as document:
        pages = len(document)
        actual = "\n".join(
            document[idx].get_textpage().get_text_bounded() for idx in range(pages)
        )
    return pages, actual


@pytest.mark.parametrize("kind", ("docx", "xlsx", "pptx", "odt", "ods", "odp"))
def test_release_richer_office_engine_to_pdf(kind, soffice, tmp_path, monkeypatch):
    """End-to-end document recognition, guarded publish, readable PDF, cleanup."""
    inputs = tmp_path / "input"
    inputs.mkdir()
    output = tmp_path / "output"
    output.mkdir()
    private_tmp = tmp_path / "qa-tmp"
    private_tmp.mkdir()
    monkeypatch.setenv("TEMP", str(private_tmp))
    monkeypatch.setenv("TMP", str(private_tmp))
    source = inputs / f"Raport_București_東京.{kind}"
    family, markers, min_pages = make_document(source)
    source_sha = _hash(source)
    assert inspect_office(source) == family
    assert inspect_source(source).kind == "office"
    assert get_targets(source, office_path=soffice) == ["pdf"]
    prior = output / f"{source.stem}.pdf"
    prior.write_bytes(b"EXISTING CUSTOMER FILE - NEVER REPLACE")
    prior_hash = _hash(prior)
    progress = []
    converted = convert_file(
        source, "pdf", output, office_path=soffice, progress=progress.append,
    )
    assert converted.name == f"{source.stem} (1).pdf"
    assert converted.is_file() and converted.stat().st_size > 500
    assert converted.read_bytes().startswith(b"%PDF-")
    assert b"%%EOF" in converted.read_bytes()[-128:]
    pages, text = _extract_pdf(converted)
    assert pages >= min_pages, f"{kind}: expected >= {min_pages}, got {pages}"
    for marker in markers:
        assert marker in text, f"{kind} missing rendered marker: {marker!r}; text={text!r}"
    assert _hash(source) == source_sha
    assert _hash(prior) == prior_hash
    assert sorted(output.iterdir()) == sorted((prior, converted))
    assert not list(tmp_path.rglob("rightconvert-office-*"))
    assert not list(tmp_path.rglob(".rightconvert-stage-*"))
    assert not list(private_tmp.rglob("*"))
    assert progress and progress[0] == 0.0 and progress[-1] == 1.0
    print(f"STAGE7_NATIVE_{kind.upper()} PASS pages={pages} bytes={converted.stat().st_size}")


def _unsafe_document(path: Path, payload: str) -> None:
    """Start with valid passive sample and add an actual active-content vector."""
    make_document(path)
    if payload == "macro":
        _alter_archive(path, {"word/vbaProject.bin": b"QA inert macro sentinel"})
    elif payload == "embedded":
        _alter_archive(path, {"ppt/embeddings/oleObject1.bin": b"QA inert OLE sentinel"})
    elif payload == "formula":
        # Works in the XLSX worksheet's real SpreadsheetML namespace.
        with zipfile.ZipFile(path) as archive:
            original = archive.read("xl/worksheets/sheet1.xml")
        formula = b'<f>WEBSERVICE("http://127.0.0.1:16666/qa")</f>'
        changed = original.replace(b"</c>", formula + b"</c>", 1)
        assert changed != original
        _alter_archive(path, {"xl/worksheets/sheet1.xml": changed})
    elif payload == "plain_formula":
        with zipfile.ZipFile(path) as archive:
            original = archive.read("xl/worksheets/sheet1.xml")
        changed = original.replace(b"</c>", b"<f>SUM(B2:B3)</f></c>", 1)
        assert changed != original
        _alter_archive(path, {"xl/worksheets/sheet1.xml": changed})
    elif payload == "external":
        # Direct local package relationship. No network lookup is performed.
        with zipfile.ZipFile(path) as archive:
            original = archive.read("ppt/_rels/presentation.xml.rels")
        changed = original.replace(
            b"</Relationships>",
            (f'<Relationship Id="external" Type="{_R}/hyperlink" '
             'Target="https://example.invalid/qa" TargetMode="External"/>'
             '</Relationships>').encode(),
        )
        _alter_archive(path, {"ppt/_rels/presentation.xml.rels": changed})
    elif payload == "odf_external":
        with zipfile.ZipFile(path) as archive:
            original = archive.read("content.xml")
        changed = original.replace(
            b"</office:text>",
            (f'<text:p xmlns:xlink="http://www.w3.org/1999/xlink"><text:a '
             'xlink:href="https://example.invalid/qa">network</text:a>'
             '</text:p></office:text>').encode(),
        )
        _alter_archive(path, {"content.xml": changed})
    else:
        raise AssertionError(payload)


@pytest.mark.parametrize("kind,payload", (
    ("docx", "macro"),
    ("pptx", "embedded"),
    ("xlsx", "formula"),
    ("xlsx", "plain_formula"),
    ("pptx", "external"),
    ("odt", "odf_external"),
))
def test_release_rejects_active_content_without_artifacts(
    kind, payload, tmp_path,
):
    source = tmp_path / f"danger-{payload}.{kind}"
    _unsafe_document(source, payload)
    original_hash = _hash(source)
    destination = tmp_path / "output"
    destination.mkdir()
    sentinel = destination / f"{source.stem}.pdf"
    sentinel.write_bytes(b"existing output")
    sentinel_hash = _hash(sentinel)
    assert inspect_office(source) is None, f"Unsafe {payload} was accepted"
    assert get_targets(source, office_path=str(QA_SOFFICE)) == []
    with pytest.raises(ConversionError):
        convert_file(source, "pdf", destination, office_path=str(QA_SOFFICE))
    assert _hash(source) == original_hash and _hash(sentinel) == sentinel_hash
    assert list(destination.iterdir()) == [sentinel]
    assert not list(tmp_path.rglob("rightconvert-office-*"))
    assert not list(tmp_path.rglob(".rightconvert-stage-*"))


def test_release_missing_libreoffice_fails_without_output(tmp_path, monkeypatch):
    source = tmp_path / "offline.docx"
    make_document(source)
    assert inspect_office(source) == "docx"
    before = _hash(source)
    missing = tmp_path / "missing" / "soffice.com"
    assert find_libreoffice(str(missing)) is None
    assert get_targets(source, office_path=str(missing)) == []
    target = tmp_path / "destination"
    with pytest.raises(ConversionError):
        convert_file(source, "pdf", target, office_path=str(missing))
    assert _hash(source) == before
    assert not target.exists()
    monkeypatch.setenv("PATH", "")
    assert find_libreoffice() is None
    assert get_targets(source) == []


def test_release_configured_binary_does_not_fall_back_to_cwd(tmp_path, monkeypatch):
    directory = tmp_path / "bin"
    directory.mkdir()
    bogus = directory / ("soffice.com" if os.name == "nt" else "soffice")
    bogus.write_bytes(b"not a program")
    monkeypatch.chdir(directory)
    monkeypatch.setenv("PATH", ".")
    assert find_libreoffice() is None
    assert find_libreoffice(bogus.name) is None
